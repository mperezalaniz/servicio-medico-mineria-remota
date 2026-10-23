"""Arma plan/Plan_100_dias_Servicio_Medico.xlsx desde plan/contenido.py (editable, con fórmulas)."""
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(__file__))
from contenido import ACTA, DMAIC, FASES, KPIS, RACI, RIESGOS, ROLES, TAREAS  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "plan", "Plan_100_dias_Servicio_Medico.xlsx")
NAVY, BLUE, PALE, ORANGE, GRID = "0E1726", "0078D4", "DCEEF8", "E66C37", "E1E1E1"
F = "Segoe UI"
th = Side(style="thin", color=GRID)
BOX = Border(left=th, right=th, top=th, bottom=th)
H = dict(font=Font(name=F, bold=True, color="FFFFFF", size=10), fill=PatternFill("solid", fgColor=BLUE),
         alignment=Alignment(vertical="center", wrap_text=True))
INPUT = Font(name=F, color="0000FF", size=10)
NORMAL = Font(name=F, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")


def head(ws, row, cols):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        for k, v in H.items():
            setattr(cell, k, v)
        cell.border = BOX
    ws.row_dimensions[row].height = 30


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=F, bold=True, size=14, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=F, size=9, color="4B5563")


def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ---------------------------------------------------------------- Cómo usar
ws = wb.active
ws.title = "Cómo usar"
title(ws, "Plan de 100 días · Team Lead de Servicios Médicos", "Emiliano Pérez Alaniz · Escenario de demostración con supuestos y datos sintéticos, no datos de una empresa")
rows = [
    ("Hojas", "Acta, Cronograma (Gantt), RACI, Riesgos, KPI y DMAIC."),
    ("Qué editar", "Solo las celdas con texto azul: inicio y duración de tareas, probabilidad e impacto de riesgos. El resto se calcula solo."),
    ("Cronograma", "Fin = Inicio + Duración - 1. Las barras del Gantt se pintan solas por semana."),
    ("Riesgos", "Puntaje = Probabilidad x Impacto (1 a 5). Alto desde 15, Medio desde 8."),
    ("Origen de las metas", "Respuesta: modelo Monte Carlo del repo (modelo/hora_de_oro.py). Exámenes, brigadistas y stock: datos sintéticos del tablero (tablero/generar_datos.py)."),
]
for i, (a, b_) in enumerate(rows, 4):
    ws.cell(row=i, column=1, value=a).font = Font(name=F, bold=True, size=10)
    c = ws.cell(row=i, column=2, value=b_)
    c.font, c.alignment = NORMAL, WRAP
widths(ws, [22, 100])

# ---------------------------------------------------------------- Acta
ws = wb.create_sheet("Acta")
title(ws, "Acta del proyecto", ACTA["nombre"])
r = 4
for lab, val in [("Propósito", [ACTA["proposito"]]), ("Objetivos al día 100", ACTA["objetivos"]), ("Alcance", ACTA["alcance"]),
                 ("Fuera de alcance", ACTA["fuera_alcance"]), ("Restricciones", ACTA["restricciones"])]:
    ws.cell(row=r, column=1, value=lab).font = Font(name=F, bold=True, size=10, color=NAVY)
    for v in val:
        c = ws.cell(row=r, column=2, value=v)
        c.font, c.alignment = NORMAL, WRAP
        r += 1
    r += 1
ws.cell(row=r, column=1, value="Fases").font = Font(name=F, bold=True, size=10, color=NAVY)
r += 1
head(ws, r, ["Días", "Fase", "Etapa DMAIC", "Qué se hace"])
for f in FASES:
    r += 1
    for j, v in enumerate(f, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.alignment, c.border = NORMAL, WRAP, BOX
widths(ws, [22, 70, 18, 80])

# ---------------------------------------------------------------- Cronograma
ws = wb.create_sheet("Cronograma")
title(ws, "Cronograma y Gantt", "Editá Inicio y Duración (texto azul). Fin y barras se calculan.")
SEM = 15
cols = ["Tarea", "Fase", "Responsable", "Entregable", "Inicio (día)", "Duración (días)", "Fin (día)"] + [f"S{k}" for k in range(1, SEM + 1)]
head(ws, 4, cols)
# fila auxiliar con el día de inicio y fin de cada semana
ws.cell(row=3, column=7, value="Semana: día inicio / fin →").font = Font(name=F, size=8, color="5A6475")
for k in range(SEM):
    col = 8 + k
    ws.cell(row=3, column=col, value=f"{k * 7 + 1}-{min(100, k * 7 + 7)}").font = Font(name=F, size=7, color="5A6475")
first = 5
for i, (t, fase, ini, dur, resp, ent) in enumerate(TAREAS):
    rr = first + i
    for j, v in enumerate([t, fase, resp, ent], 1):
        c = ws.cell(row=rr, column=j, value=v)
        c.font, c.alignment, c.border = NORMAL, WRAP, BOX
    ws.cell(row=rr, column=5, value=ini).font = INPUT
    ws.cell(row=rr, column=6, value=dur).font = INPUT
    ws.cell(row=rr, column=7, value=f"=E{rr}+F{rr}-1").font = NORMAL
    for col in range(5, 8 + SEM):
        ws.cell(row=rr, column=col).border = BOX
        ws.cell(row=rr, column=col).alignment = Alignment(horizontal="center", vertical="center")
last = first + len(TAREAS) - 1
rng_ = f"H{first}:{get_column_letter(7 + SEM)}{last}"
# semana k (columna H = k0): días 7k+1 a 7k+7
ws.conditional_formatting.add(rng_, FormulaRule(
    formula=[f"AND($E{first}<=(COLUMN()-8)*7+7,$G{first}>=(COLUMN()-8)*7+1)"], fill=PatternFill("solid", fgColor=BLUE)))
ws.cell(row=last + 2, column=1, value="Duración total del plan (días)").font = Font(name=F, bold=True, size=10)
ws.cell(row=last + 2, column=7, value=f"=MAX(G{first}:G{last})").font = Font(name=F, bold=True, size=10)
ws.cell(row=last + 3, column=1, value="Tareas").font = Font(name=F, bold=True, size=10)
ws.cell(row=last + 3, column=7, value=f"=COUNTA(A{first}:A{last})").font = Font(name=F, bold=True, size=10)
widths(ws, [46, 22, 26, 36, 9, 9, 9] + [4.2] * SEM)
ws.freeze_panes = "E5"

# ---------------------------------------------------------------- RACI
ws = wb.create_sheet("RACI")
title(ws, "Matriz RACI", "R: hace · A: responde por el resultado · C: se consulta · I: se informa")
head(ws, 4, ["Actividad"] + ROLES + ["Control: una sola A"])
for i, (act, letras) in enumerate(RACI, 5):
    ws.cell(row=i, column=1, value=act).font = NORMAL
    for j, l in enumerate(letras, 2):
        c = ws.cell(row=i, column=j, value=l or None)
        c.font, c.alignment, c.border = Font(name=F, size=10, bold=l == "A"), Alignment(horizontal="center"), BOX
    end = get_column_letter(1 + len(ROLES))
    ws.cell(row=i, column=2 + len(ROLES), value=f'=IF(COUNTIF(B{i}:{end}{i},"A")=1,"OK","Revisar")').font = NORMAL
    ws.cell(row=i, column=1).border = BOX
ws.conditional_formatting.add(f"B5:{get_column_letter(1 + len(ROLES))}{4 + len(RACI)}",
                              CellIsRule(operator="equal", formula=['"A"'], fill=PatternFill("solid", fgColor=PALE)))
widths(ws, [34] + [14] * len(ROLES) + [16])

# ---------------------------------------------------------------- Riesgos
ws = wb.create_sheet("Riesgos")
title(ws, "Registro de riesgos", "Probabilidad e impacto de 1 a 5 (texto azul, editable)")
head(ws, 4, ["Riesgo", "Probabilidad", "Impacto", "Puntaje", "Nivel", "Respuesta"])
for i, (rg, p, im, resp) in enumerate(RIESGOS, 5):
    ws.cell(row=i, column=1, value=rg).font = NORMAL
    ws.cell(row=i, column=2, value=p).font = INPUT
    ws.cell(row=i, column=3, value=im).font = INPUT
    ws.cell(row=i, column=4, value=f"=B{i}*C{i}").font = NORMAL
    ws.cell(row=i, column=5, value=f'=IF(D{i}>=15,"Alto",IF(D{i}>=8,"Medio","Bajo"))').font = NORMAL
    ws.cell(row=i, column=6, value=resp).font = NORMAL
    for j in range(1, 7):
        ws.cell(row=i, column=j).border = BOX
        ws.cell(row=i, column=j).alignment = WRAP if j in (1, 6) else Alignment(horizontal="center", vertical="top")
lr = 4 + len(RIESGOS)
ws.conditional_formatting.add(f"E5:E{lr}", CellIsRule(operator="equal", formula=['"Alto"'], fill=PatternFill("solid", fgColor="FDE7E9")))
ws.conditional_formatting.add(f"E5:E{lr}", CellIsRule(operator="equal", formula=['"Medio"'], fill=PatternFill("solid", fgColor="FBE3D6")))
ws.cell(row=lr + 2, column=1, value="Riesgos de nivel alto").font = Font(name=F, bold=True, size=10)
ws.cell(row=lr + 2, column=4, value=f'=COUNTIF(E5:E{lr},"Alto")').font = Font(name=F, bold=True, size=10)
widths(ws, [46, 12, 10, 10, 10, 70])

# ---------------------------------------------------------------- KPI
ws = wb.create_sheet("KPI")
title(ws, "Árbol de indicadores", "Línea base: Modelo = modelo Monte Carlo del repo · Sintético = datos de demostración · A medir = se toma en los primeros 30 días")
head(ws, 4, ["Dimensión", "Indicador", "Cómo se calcula", "Línea base", "Meta día 100", "Frecuencia", "Origen de la línea base"])
for i, k in enumerate(KPIS, 5):
    for j, v in enumerate(k, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.font, c.alignment, c.border = NORMAL, WRAP, BOX
widths(ws, [13, 32, 52, 18, 16, 20, 18])

# ---------------------------------------------------------------- DMAIC
ws = wb.create_sheet("DMAIC")
title(ws, "Dos mejoras con DMAIC")
r = 3
for d in DMAIC:
    r += 1
    ws.cell(row=r, column=1, value=d["titulo"]).font = Font(name=F, bold=True, size=12, color=NAVY)
    for etapa, nombre in [("D", "Definir"), ("M", "Medir"), ("A", "Analizar"), ("I", "Mejorar"), ("C", "Controlar")]:
        r += 1
        a = ws.cell(row=r, column=1, value=nombre)
        a.font, a.fill, a.border = Font(name=F, bold=True, size=10, color="FFFFFF"), PatternFill("solid", fgColor=ORANGE if etapa == "I" else BLUE), BOX
        a.alignment = Alignment(vertical="top")
        c = ws.cell(row=r, column=2, value=d[etapa])
        c.font, c.alignment, c.border = NORMAL, WRAP, BOX
        ws.row_dimensions[r].height = 46
    r += 1
widths(ws, [14, 120])
ws["A1"].comment = Comment("Cifras de respuesta: modelo/resultados.json. Cifras de stock: tablero/tablero.json (sintético).", "Emiliano Pérez Alaniz")

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
wb.save(OUT)
print(OUT)
